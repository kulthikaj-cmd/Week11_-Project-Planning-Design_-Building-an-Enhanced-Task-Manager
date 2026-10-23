# src/word_processor.py
import os
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from utils import setup_logging

logger = setup_logging(__name__)

class WordProcessor:
    def __init__(self):
        logger.info("WordProcessor initialized.")

    def create_document(self):
        return docx.Document()

    def add_heading(self, doc, text, level=1):
        doc.add_heading(text, level=level)

    def add_paragraph(self, doc, text, bold=False, italic=False, align='left'):
        p = doc.add_paragraph(text)
        if p.runs:
            p.runs[0].bold = bold
            p.runs[0].italic = italic
        if align == 'center':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return p

    def add_table(self, doc, data, col_widths=None):
        table = doc.add_table(rows=len(data), cols=len(data[0]))
        table.style = 'Table Grid'
        if col_widths:
            for i, width in enumerate(col_widths):
                table.columns[i].width = Inches(width)
        for r_idx, row in enumerate(data):
            for c_idx, val in enumerate(row):
                cell = table.rows[r_idx].cells[c_idx]
                cell.text = str(val)
                if r_idx == 0 and cell.paragraphs[0].runs:
                    cell.paragraphs[0].runs[0].bold = True

    def save_document(self, doc, output_path):
        doc.save(output_path)
        logger.info(f"Word document saved to '{output_path}'.")
